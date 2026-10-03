#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
Voegt de hand-dia's toe aan het Werkcollege 8-deck.

    python3 powerpoints/wc8_handen.py --ronde 6      (tekent de platen)
    python3 powerpoints/vul_wc8_aan.py               (zet ze in het deck)

Het origineel wordt NOOIT overschreven. Data_Science_Werkcollege8.pptx is op
30 september door PowerPoint opgeslagen -- er is met de hand aan gewerkt, en
maak_presentatie_wc8.py opnieuw draaien zou dat weggooien. Dit script leest het
bestaande deck en schrijft een nieuwe kopie ernaast.

De dia's gebruiken dezelfde stijl als de rest van het deck (dekstijl.py), zodat
ze er niet uit springen.
"""
import os
import sys

from pptx import Presentation
from pptx.util import Inches, Pt

HIER = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, HIER)

from dekstijl import ACCENT, GEDEMPT, PRIMAIR, kaart, kop, tekstblok  # noqa: E402

PLOTS = os.path.join(HIER, "plots_wc8")
BRON = os.path.join(HIER, "Data_Science_Werkcollege8.pptx")
UITVOER = os.path.join(HIER, "Data_Science_Werkcollege8_met_handen.pptx")

# (plaat, eyebrow, titel, regels onder de plaat, sprekersnotitie)
#
# De platen komen uit wc8_handen.py. Welke hand welke plaat is, staat in de
# bestandsnaam: s = simulatie, t = tafel, h = handnummer.
DIAS = [
    ("hand_winst_r6_s15t3h1.png",
     "EEN HAND VAN DICHTBIJ",
     "De grootste pot van het toernooi",
     [("9♥9♦ callt preflop, krijgt een grote raise van A♦A♣ over zich heen — en callt "
       "alsnog. Dan komt de flop: 5♣ J♥ 9♠.", 15, False, PRIMAIR, 10),
      ("Dat is een set negens, en vanaf dat moment raisen vier spelers elkaar de stack "
       "uit. Drie van hen verliezen alles.", 15, False, PRIMAIR, 10),
      ("Let op wat hier niet gebeurt: niemand rekent uit dat A♦A♣ ná die flop niet meer "
       "de beste hand is. Je bot kijkt naar zijn eigen kaarten, niet naar het bord.",
       14, False, GEDEMPT, 0)],
     """Dit is de hand om mee te openen: hij is spectaculair en hij maakt een punt.

     Vraag aan de zaal: wie van deze zes had moeten folden, en wanneer?
     Het antwoord is A♦A♣ op de flop. Preflop was die hand de beste van de tafel,
     na 5♣ J♥ 9♠ niet meer. Geen enkele bot in de klas kijkt daarnaar.

     Dat is precies het bruggetje naar de uitgebreide log: de kolom 'bord' is
     nieuw, en dit is waarom hij er is."""),

    ("hand_verlies_r6_s4t0h37.png",
     "EEN HAND VAN DICHTBIJ",
     "De duurste hand: vier keer verhogen met de op één na beste hand",
     [("Q♠J♥ flopt een paar boeren op A♥J♦3♣. Dat voelt sterk — tot je ziet dat de "
       "tegenstander K♣A♦ heeft en dus azen.", 15, False, PRIMAIR, 10),
      ("Vier raises op de flop, all-in op de turn, 2.200 chips weg. In één hand, van de "
       "3.000 waarmee deze bot aan die tafel zat.", 15, False, PRIMAIR, 10),
      ("Een aas op het bord is gevaarlijk juist omdat iedereen graag met een aas speelt. "
       "Dat kun je in de data terugzien.", 14, False, GEDEMPT, 0)],
     """Hier gaat het om het verschil tussen 'sterk' en 'de beste'.

     Een paar boeren is objectief een nette hand. Maar het bord heeft een aas, en
     de helft van de klas speelt elke hand met een aas. De kans dat iemand die nog
     meedoet er een heeft is dus groot.

     Dit is te meten met de uitgebreide log: hoe vaak wordt er doorbetaald op een
     bord met een aas, en wat levert dat gemiddeld op? Goede opdracht voor Deel 4."""),

    ("hand_verlies_r6_s2t3h5.png",
     "EEN HAND VAN DICHTBIJ",
     "Hetzelfde patroon, tegen een referentiebot",
     [("J♠Q♣ op K♣A♥Q♥: opnieuw een paar, opnieuw niet het hoogste. De turn brengt K♦ "
       "en dan heeft Referentie_Bluffer twee paar azen en heren.", 15, False, PRIMAIR, 10),
      ("Dezelfde student verliest hier 2.010. Dat is geen pech: van alle 27 bots in de "
       "klas heeft deze de grootste spreiding per hand (238 tegen een mediaan van 131).",
       15, False, PRIMAIR, 10),
      ("Hij won óók de grootste pot van het toernooi — en eindigde op 934, onder zijn "
       "startstack.", 14, False, ACCENT, 0)],
     """Dit is de dia waar de les in zit.

     Dezelfde bot staat in alle drie de uitschieters: de grootste pot gewonnen én
     de twee duurste handen verloren. Dat is geen toeval, het is spreiding, en die
     is te meten.

     Maak de link naar Deel 5: één toernooi is twintig simulaties, en een bot met
     deze spreiding kan op basis van geluk zomaar vijf plekken verschuiven.
     Agressiever spelen maakt je niet beter, het maakt je uitslag onzekerder."""),
]


def voeg_dia_toe(prs, plaat, eyebrow, titel, regels, notitie):
    dia = prs.slides.add_slide(prs.slide_layouts[6])
    # De achtergrond en de kop komen uit dekstijl, net als bij de bestaande dia's.
    from dekstijl import BG
    from pptx.enum.shapes import MSO_SHAPE
    vlak = dia.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, prs.slide_width, prs.slide_height)
    vlak.fill.solid()
    vlak.fill.fore_color.rgb = BG
    vlak.line.fill.background()
    vlak.shadow.inherit = False
    kop(dia, eyebrow, titel)

    pad = os.path.join(PLOTS, plaat)
    if not os.path.exists(pad):
        print(f"   ONTBREEKT: {plaat} -- draai eerst wc8_handen.py")
        return dia

    beeld = dia.shapes.add_picture(pad, 0, Inches(1.58), height=Inches(3.30))
    if beeld.width > Inches(12.2):
        beeld.height = int(beeld.height * Inches(12.2) / beeld.width)
        beeld.width = Inches(12.2)
    beeld.left = int((prs.slide_width - beeld.width) / 2)
    beeld.top = Inches(1.58)

    boven = 5.05
    kaart(dia, Inches(0.8), Inches(boven), Inches(11.7), Inches(2.25))
    tekstblok(dia, Inches(1.2), Inches(boven + 0.16), Inches(10.9), Inches(1.93), regels)

    if notitie:
        dia.notes_slide.notes_text_frame.text = "\n".join(
            regel.strip() for regel in notitie.strip().split("\n"))
    return dia


def main():
    if not os.path.exists(BRON):
        raise SystemExit(f"Bronbestand niet gevonden: {BRON}")
    prs = Presentation(BRON)
    voor = len(prs.slides._sldIdLst)
    print(f"{voor} dia's in het origineel")
    for spec in DIAS:
        voeg_dia_toe(prs, *spec)
        print(f"   + {spec[2]}")
    prs.save(UITVOER)
    na = len(prs.slides._sldIdLst)
    print(f"\n{na} dia's -> {os.path.basename(UITVOER)}")
    print(f"Het origineel ({os.path.basename(BRON)}) is niet aangeraakt.")
    return 0


if __name__ == "__main__":
    sys.exit(main())
